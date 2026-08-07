# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import io
import re

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    openpyxl = None


def _clean_int(val):
    if val is None or val is False:
        return None
    try:
        f = float(str(val).strip())
        if f.is_integer():
            return int(f)
    except (ValueError, TypeError):
        pass
    return None


def _clean_float(val):
    if val is None or val is False:
        return 0.0
    s = str(val).strip()
    # Remove currency symbols and spaces
    s = re.sub(r'[$€£\s]', '', s)
    # Handle European format: 1.500,00 -> 1500.00
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    # Remove anything that is not digit or dot or minus
    s = re.sub(r'[^\d.\-]', '', s)
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def _clean_str(val):
    if val is None or val is False:
        return False
    s = str(val).strip()
    if not s:
        return False
    # Remove trailing .0 added by Excel on integer-like numbers
    if s.endswith('.0') and s[:-2].replace('.', '').replace('-', '').isdigit():
        s = s[:-2]
    return s


class SmartExcelUpdate(models.TransientModel):
    _name = 'smart.excel.update'
    _description = 'Smart Excel Mass Update & Creation Wizard'

    state = fields.Selection([
        ('draft', 'Draft'),
        ('done', 'Completed'),
    ], string='Status', default='draft')

    # main_mode kept for routing view sections
    main_mode = fields.Selection([
        ('export', 'Export to Excel'),
        ('import', 'Import to Odoo'),
    ], string='Operation', default='export', required=True)

    # Products populated automatically from list selection (export only)
    product_ids = fields.Many2many(
        'product.template',
        string='Selected Products',
    )

    exported_file = fields.Binary('Exported File', attachment=False)
    import_file = fields.Binary('Select Excel File (.xlsx)', attachment=False)
    file_name = fields.Char('File Name')
    log_message = fields.Text('Result Log', readonly=True)

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        active_ids = self.env.context.get('active_ids', [])
        active_model = self.env.context.get('active_model')
        if active_model == 'product.template' and active_ids:
            res['product_ids'] = [(6, 0, active_ids)]
        return res

    # ------------------------------------------------------------------
    # EXPORT
    # ------------------------------------------------------------------
    def action_generate_template(self):
        """Export selected products to Excel with their Odoo IDs."""
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("The 'openpyxl' Python library is required. Contact your system administrator."))

        user_lang = self.env.user.lang or 'en_US'
        is_spanish = user_lang.startswith('es')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Productos" if is_spanish else "Products"

        # ---- Header styling ----
        header_fill = PatternFill(start_color="0F1729", end_color="0F1729", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center = Alignment(horizontal='center', vertical='center')

        if is_spanish:
            headers = ["ID", "Ref. Interna", "Código de Barras", "Nombre",
                       "Precio de Venta", "Costo", "Categoría", "Peso (kg)", "Descripción"]
        else:
            headers = ["ID", "Internal Reference", "Barcode", "Name",
                       "Sales Price", "Cost", "Category", "Weight (kg)", "Description"]

        col_widths = [12, 20, 20, 35, 15, 15, 25, 12, 40]
        for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            ws.column_dimensions[chr(64 + col_idx)].width = width

        ws.row_dimensions[1].height = 22

        # ---- Data rows ----
        products = self.product_ids
        if not products:
            raise UserError(
                _("No products selected. Please go to Inventory > Products, select the products you want to export, then use the Action menu.")
                if not is_spanish else
                _("No hay productos seleccionados. Ve a Inventario > Productos, selecciona los productos que quieres exportar y usa el menú Acción.")
            )

        for row_num, product in enumerate(products, 2):
            ws.cell(row=row_num, column=1, value=product.id)
            ws.cell(row=row_num, column=2, value=product.default_code or '')
            ws.cell(row=row_num, column=3, value=product.barcode or '')
            ws.cell(row=row_num, column=4, value=product.name or '')
            ws.cell(row=row_num, column=5, value=product.list_price or 0.0)
            ws.cell(row=row_num, column=6, value=product.standard_price or 0.0)
            ws.cell(row=row_num, column=7, value=product.categ_id.display_name or '')
            ws.cell(row=row_num, column=8, value=product.weight or 0.0)
            ws.cell(row=row_num, column=9, value=product.description_sale or '')

        count_exported = len(products)
        filename = "Productos_Exportados.xlsx" if is_spanish else "Exported_Products.xlsx"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_data = base64.b64encode(output.read())

        if is_spanish:
            log_msg = (
                f"Exportacion Completada con Exito!\n"
                f"----------------------------------------\n"
                f"Productos Exportados: {count_exported}\n"
                f"Archivo: {filename}\n\n"
                f"Haz clic en 'Descargar Archivo Excel' para guardar."
            )
        else:
            log_msg = (
                f"Export Completed Successfully!\n"
                f"----------------------------------------\n"
                f"Products Exported: {count_exported}\n"
                f"File: {filename}\n\n"
                f"Click 'Download Excel File' to save."
            )

        self.write({
            'exported_file': file_data,
            'file_name': filename,
            'log_message': log_msg,
            'state': 'done',
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'smart.excel.update',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_download_file(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': (
                f'/web/content/?model=smart.excel.update'
                f'&id={self.id}'
                f'&field=exported_file'
                f'&filename={self.file_name}'
                f'&download=true'
            ),
            'target': 'self',
        }

    # ------------------------------------------------------------------
    # IMPORT
    # ------------------------------------------------------------------
    def action_import_data(self):
        """
        Import/sync an Excel file back to Odoo.
        - Rows WITH an ID in column A  -> update that product.template record.
        - Rows WITHOUT an ID           -> create a new product.template.
        No product selection is required to run this operation.
        """
        self.ensure_one()
        if not self.import_file:
            raise UserError(_("Please select an Excel (.xlsx) file before clicking Import."))

        if openpyxl is None:
            raise UserError(_("The 'openpyxl' Python library is required. Contact your system administrator."))

        try:
            file_data = base64.b64decode(self.import_file)
            input_stream = io.BytesIO(file_data)
            wb = openpyxl.load_workbook(input_stream, data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_("Invalid Excel spreadsheet (.xlsx): %s") % str(e))

        updated_count = 0
        created_count = 0
        errors = []

        user_lang = self.env.user.lang or 'en_US'
        is_spanish = user_lang.startswith('es')

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not row or all(v is None for v in row):
                continue

            raw_id       = _clean_int(row[0]) if len(row) > 0 else None
            default_code = _clean_str(row[1]) if len(row) > 1 else False
            barcode      = _clean_str(row[2]) if len(row) > 2 else False
            name         = _clean_str(row[3]) if len(row) > 3 else False
            list_price   = _clean_float(row[4]) if len(row) > 4 else 0.0
            std_price    = _clean_float(row[5]) if len(row) > 5 else 0.0
            categ_name   = _clean_str(row[6]) if len(row) > 6 else False
            weight       = _clean_float(row[7]) if len(row) > 7 else 0.0
            description  = _clean_str(row[8]) if len(row) > 8 else False

            vals = {}
            if default_code is not False: vals['default_code'] = default_code
            if barcode      is not False: vals['barcode']       = barcode
            if name         is not False: vals['name']          = name
            vals['list_price']      = list_price
            vals['standard_price']  = std_price
            vals['weight']          = weight
            if description  is not False: vals['description_sale'] = description

            # Category lookup by name or complete_name
            if categ_name:
                categ = self.env['product.category'].search(
                    [('name', '=', categ_name)], limit=1
                )
                if not categ:
                    categ = self.env['product.category'].search(
                        [('complete_name', '=', categ_name)], limit=1
                    )
                if categ:
                    vals['categ_id'] = categ.id

            try:
                with self.env.cr.savepoint():
                    if raw_id:
                        product = self.env['product.template'].browse(raw_id)
                        if product.exists():
                            product.write(vals)
                            updated_count += 1
                        else:
                            msg = (
                                f"Fila {row_idx}: ID {raw_id} no encontrado en Odoo."
                                if is_spanish else
                                f"Row {row_idx}: ID {raw_id} not found in Odoo."
                            )
                            errors.append(msg)
                    else:
                        if not name:
                            msg = (
                                f"Fila {row_idx}: Omitido — el campo 'Nombre' esta vacio."
                                if is_spanish else
                                f"Row {row_idx}: Skipped — 'Name' field is empty."
                            )
                            errors.append(msg)
                        else:
                            self.env['product.template'].create(vals)
                            created_count += 1
            except Exception as e:
                err_clean = str(e).split('\n')[0]
                msg = (
                    f"Fila {row_idx}: Error — {err_clean}."
                    if is_spanish else
                    f"Row {row_idx}: Error — {err_clean}."
                )
                errors.append(msg)

        if is_spanish:
            log_msg = (
                f"Proceso Finalizado con Exito!\n"
                f"----------------------------------------\n"
                f"Productos Actualizados: {updated_count}\n"
                f"Productos Nuevos Creados: {created_count}\n"
            )
            if errors:
                log_msg += f"\nFilas con Errores ({len(errors)}):\n" + "\n".join(errors)
        else:
            log_msg = (
                f"Process Finished Successfully!\n"
                f"----------------------------------------\n"
                f"Updated Existing Products: {updated_count}\n"
                f"Created New Products: {created_count}\n"
            )
            if errors:
                log_msg += f"\nIgnored Rows / Errors ({len(errors)}):\n" + "\n".join(errors)

        self.log_message = log_msg
        self.state = 'done'

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'smart.excel.update',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }