# -*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import io
import re

try:
    import openpyxl
except ImportError:
    openpyxl = None

def _clean_int(val):
    if val is None or val is False:
        return None
    try:
        f = float(str(val).strip())
        if f.is_integer():
            return int(f)
    except (ValueError, TypeError):
        pass
    return None

def _clean_float(val):
    if val is None or val is False:
        return 0.0
    s = str(val).strip().replace('$', '').replace('€', '').replace(' ', '')
    s = s.replace(',', '.')
    s = re.sub(r'[^\d.-]', '', s)
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0

def _clean_str(val):
    if val is None or val is False:
        return False
    s = str(val).strip()
    if not s:
        return False
    if s.endswith('.0') and s[:-2].replace('.', '').isdigit():
        s = s[:-2]
    return s

class SmartExcelUpdate(models.TransientModel):
    _name = 'smart.excel.update'
    _description = 'Smart Excel Mass Update & Creation Wizard'

    state = fields.Selection([
        ('draft', 'Draft'),
        ('done', 'Completed')
    ], string='Status', default='draft')

    main_mode = fields.Selection([
        ('export', 'Export to Excel'),
        ('import', 'Import to Odoo')
    ], string='Operation', default='export', required=True)
    
    export_type = fields.Selection([
        ('update', 'Export selected products with IDs (for updating existing products)'),
        ('blank', 'Export selected products with blank IDs (for creating new products)')
    ], string='Export Option', default='update')
    
    product_ids = fields.Many2many(
        'product.template',
        string='Selected Products'
    )

    exported_file = fields.Binary('Exported File', attachment=False)
    import_file = fields.Binary('Select Excel File', attachment=False)
    file_name = fields.Char('File Name')
    log_message = fields.Text('Result Log', readonly=True)

    @api.model
    def default_get(self, fields_list):
        res = super(SmartExcelUpdate, self).default_get(fields_list)
        active_ids = self.env.context.get('active_ids', [])
        active_model = self.env.context.get('active_model')
        if active_model == 'product.template' and active_ids:
            res['product_ids'] = [(6, 0, active_ids)]
        return res

    def action_generate_template(self):
        if not openpyxl:
            raise UserError(_("Please install the openpyxl python library to use this module."))
        
        user_lang = self.env.user.lang or 'en_US'
        is_spanish = user_lang.startswith('es')

        products = self.product_ids
        if not products:
            active_ids = self.env.context.get('active_ids', [])
            if active_ids:
                products = self.env['product.template'].browse(active_ids)
                
        if self.export_type == 'update' and not products:
            raise UserError(_("No products selected. Please select products from the list view first before exporting."))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products Data"
        
        if is_spanish:
            headers = ['ID (Dejar vacio para nuevos)', 'Referencia Interna', 'Codigo de Barras', 'Nombre', 'Precio de Venta', 'Costo', 'Categoria', 'Peso (kg)', 'Descripcion']
        else:
            headers = ['ID (Leave empty for new products)', 'Internal Reference', 'Barcode', 'Name', 'Sales Price', 'Cost', 'Category', 'Weight (kg)', 'Description']
            
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)
            
        if products:
            for row_num, product in enumerate(products, 2):
                ws.cell(row=row_num, column=1, value=product.id if self.export_type == 'update' else "")
                ws.cell(row=row_num, column=2, value=product.default_code or '')
                ws.cell(row=row_num, column=3, value=product.barcode or '')
                ws.cell(row=row_num, column=4, value=product.name or '')
                ws.cell(row=row_num, column=5, value=product.list_price or 0.0)
                ws.cell(row=row_num, column=6, value=product.standard_price or 0.0)
                ws.cell(row=row_num, column=7, value=product.categ_id.display_name or '')
                ws.cell(row=row_num, column=8, value=product.weight or 0.0)
                ws.cell(row=row_num, column=9, value=product.description_sale or '')
            count_exported = len(products)
        else:
            ws.cell(row=2, column=1, value="")
            ws.cell(row=2, column=2, value="REF-001")
            ws.cell(row=2, column=3, value="1234567890123")
            ws.cell(row=2, column=4, value="Producto de Ejemplo" if is_spanish else "Sample New Product")
            ws.cell(row=2, column=5, value=100.0)
            ws.cell(row=2, column=6, value=50.0)
            ws.cell(row=2, column=7, value="All")
            ws.cell(row=2, column=8, value=1.5)
            ws.cell(row=2, column=9, value="Descripción corta del producto" if is_spanish else "Short product description")
            count_exported = 1

        if self.export_type == 'update':
            filename = "Productos_Exportados.xlsx" if is_spanish else "Exported_Products.xlsx"
        else:
            filename = "Plantilla_Creacion_Productos.xlsx" if is_spanish else "Product_Creation_Template.xlsx"
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        file_data = base64.b64encode(output.read())
        
        if is_spanish:
            log_msg = f"🎉 ¡Exportación Completada con Éxito!\n"
            log_msg += f"----------------------------------------\n"
            log_msg += f"📦 Productos Exportados al Archivo: {count_exported}\n"
            log_msg += f"📄 Archivo Generado: {filename}\n\n"
            log_msg += f"👉 Haz clic en el botón de abajo para descargar tu archivo Excel."
        else:
            log_msg = f"🎉 Export Completed Successfully!\n"
            log_msg += f"----------------------------------------\n"
            log_msg += f"📦 Products Exported to File: {count_exported}\n"
            log_msg += f"📄 Generated File Name: {filename}\n\n"
            log_msg += f"👉 Click the download button below to save your Excel file."

        self.write({
            'exported_file': file_data,
            'file_name': filename,
            'log_message': log_msg,
            'state': 'done',
        })
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'smart.excel.update',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_download_file(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model=smart.excel.update&id={self.id}&field=exported_file&filename={self.file_name}&download=true',
            'target': 'self',
        }

    def action_import_data(self):
        if not self.import_file:
            raise UserError(_("Please select an Excel file from your computer before clicking Import."))
            
        try:
            file_data = base64.b64decode(self.import_file)
            input_stream = io.BytesIO(file_data)
            wb = openpyxl.load_workbook(input_stream, data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_("Invalid Excel spreadsheet (.xlsx): %s") % str(e))
            
        updated_count = 0
        created_count = 0
        errors = []
        
        user_lang = self.env.user.lang or 'en_US'
        is_spanish = user_lang.startswith('es')

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not row or all(v is None for v in row):
                continue
                
            raw_id = _clean_int(row[0]) if len(row) > 0 else None
            default_code = _clean_str(row[1]) if len(row) > 1 else False
            barcode = _clean_str(row[2]) if len(row) > 2 else False
            name = _clean_str(row[3]) if len(row) > 3 else False
            list_price = _clean_float(row[4]) if len(row) > 4 else 0.0
            standard_price = _clean_float(row[5]) if len(row) > 5 else 0.0
            categ_name = _clean_str(row[6]) if len(row) > 6 else False
            weight = _clean_float(row[7]) if len(row) > 7 else 0.0
            description = _clean_str(row[8]) if len(row) > 8 else False
            
            vals = {}
            if default_code is not False: vals['default_code'] = default_code
            if barcode is not False: vals['barcode'] = barcode
            if name is not False: vals['name'] = name
            vals['list_price'] = list_price
            vals['standard_price'] = standard_price
            vals['weight'] = weight
            if description is not False: vals['description_sale'] = description
            
            if categ_name:
                categ = self.env['product.category'].search([('name', '=', categ_name)], limit=1)
                if not categ:
                    categ = self.env['product.category'].search([('complete_name', '=', categ_name)], limit=1)
                if categ:
                    vals['categ_id'] = categ.id

            try:
                with self.env.cr.savepoint():
                    if raw_id:
                        product = self.env['product.template'].browse(raw_id)
                        if product.exists():
                            product.write(vals)
                            updated_count += 1
                        else:
                            msg = f"Fila {row_idx}: Producto ID {raw_id} no encontrado en Odoo." if is_spanish else f"Row {row_idx}: Product ID {raw_id} not found in Odoo."
                            errors.append(msg)
                    else:
                        if not name:
                            msg = f"Fila {row_idx}: Omitido producto nuevo porque el 'Nombre' está vacío." if is_spanish else f"Row {row_idx}: Skipped new product because 'Name' is empty."
                            errors.append(msg)
                        else:
                            self.env['product.template'].create(vals)
                            created_count += 1
            except Exception as e:
                err_clean = str(e).split('\n')[0]
                msg = f"Fila {row_idx}: Error ({err_clean})." if is_spanish else f"Row {row_idx}: Error ({err_clean})."
                errors.append(msg)
                
        if is_spanish:
            log_msg = f"🎉 ¡Proceso Finalizado con Éxito!\n"
            log_msg += f"----------------------------------------\n"
            log_msg += f"✏️ Productos Existentes Actualizados: {updated_count}\n"
            log_msg += f"✨ Productos Nuevos Creados: {created_count}\n"
            if errors:
                log_msg += f"\n⚠️ Filas Ignoradas / Errores ({len(errors)}):\n" + "\n".join(errors)
        else:
            log_msg = f"🎉 Process Finished Successfully!\n"
            log_msg += f"----------------------------------------\n"
            log_msg += f"✏️ Updated Existing Products: {updated_count}\n"
            log_msg += f"✨ Created New Products: {created_count}\n"
            if errors:
                log_msg += f"\n⚠️ Ignored Rows / Errors ({len(errors)}):\n" + "\n".join(errors)
            
        self.log_message = log_msg
        self.state = 'done'
        
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'smart.excel.update',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }